import pymupdf
from docx import Document
from pathlib import Path


def extract_resume(file_path: str) -> str:
    """
    Extract text from PDF or DOCX resume.
    """

    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    extension = path.suffix.lower()

    text = ""

    if extension == ".pdf":
        pdf = pymupdf.open(path)

        try:
            for page in pdf:
                text += page.get_text() + "\n"
        finally:
            pdf.close()

    elif extension == ".docx":
        document = Document(path)

        for paragraph in document.paragraphs:
            text += paragraph.text + "\n"

        for table in document.tables:
            for row in table.rows:
                text += " | ".join(
                    cell.text for cell in row.cells
                ) + "\n"

    else:
        raise ValueError(
            "Unsupported file format. Upload PDF or DOCX."
        )

    text = text.strip()

    if not text:
        raise ValueError(
            "No readable text found in the resume."
        )

    return text
